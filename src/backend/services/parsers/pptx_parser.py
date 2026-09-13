"""PPTX 解析器：python-pptx 逐页解析，``page_no`` = 幻灯片序号（1 起）。"""

from __future__ import annotations

import logging
from pathlib import Path

from .base import Block, make_anchor
from ...utils.textutil import clean_text

logger = logging.getLogger(__name__)

__all__ = ["parse"]


def _shape_text(shape: object) -> list[str]:
    """提取单个 shape 的文本（文本框 / 表格 / 组合）。"""
    texts: list[str] = []
    try:
        if getattr(shape, "has_text_frame", False):
            value = clean_text(shape.text_frame.text or "")  # type: ignore[attr-defined]
            if value:
                texts.append(value)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:  # type: ignore[attr-defined]
                cells = [clean_text(cell.text or "") for cell in row.cells]
                row_text = " | ".join(c for c in cells if c)
                if row_text:
                    texts.append(row_text)
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            for child in shape.shapes:  # type: ignore[attr-defined]
                texts.extend(_shape_text(child))
    except Exception:  # noqa: BLE001 - 单个 shape 失败不阻断
        pass
    return texts


def parse(path: Path, **_: object) -> list[Block]:
    """逐幻灯片解析 PPTX。

    Args:
        path: PPTX 文件路径。
        **_: 兼容统一签名。

    Returns:
        ``Block`` 列表，每张幻灯片一个块（有文本才产生），``page_no`` = 幻灯片序号。

    Raises:
        RuntimeError: 文件打开失败。
    """
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"PPTX 打开失败：{type(exc).__name__}") from exc

    blocks: list[Block] = []
    for index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            texts.extend(_shape_text(shape))
        page_text = clean_text("\n".join(texts))
        if not page_text:
            continue
        section = f"第 {index} 页幻灯片"
        blocks.append(
            Block(
                text=page_text,
                page_no=index,
                section=section,
                anchor=make_anchor(page_text, slide=index),
                order=len(blocks),
            )
        )

    logger.info("PPTX 解析完成", extra={"extra_fields": {"slides": len(blocks)}})
    return blocks


def page_count(path: Path) -> int:
    """返回幻灯片总数。"""
    from pptx import Presentation

    try:
        return len(Presentation(str(path)).slides)
    except Exception:  # noqa: BLE001
        return 0
