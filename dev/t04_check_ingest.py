"""T04 自检：真实文件解析入库。

生成夹具（放到 dev/_out/）并全部走「上传 → 解析 → status=ready」：
    * 多页 PDF（手工构造，pypdf 可提取文本）——验证切片带正确 page_no 整数；
    * 扫描版 PDF（Pillow 图片转 PDF，无文本层）——验证 warning 提示；
    * DOCX、PPTX（≥3 页）、MD、TXT；
    * 删除文档后 chunks 归零。

运行前启动服务：
    PYTHONPATH=src .venv/Scripts/python.exe dev/t04_check_ingest.py
"""

from __future__ import annotations

import io
import sqlite3
import sys
import time
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dev" / "_out"
DB_PATH = ROOT / "data" / "mrnibble.db"
BASE = "http://127.0.0.1:8760"
PASSED = 0


def check(name: str, condition: bool, extra: str = "") -> None:
    """断言并记录。"""
    global PASSED
    mark = "PASS" if condition else "FAIL"
    print(f"[{mark}] {name}" + (f"  -> {extra}" if extra else ""))
    if not condition:
        raise SystemExit(f"断言失败：{name} {extra}")
    PASSED += 1


# ── 夹具生成 ────────────────────────────────────────────
def _pdf_escape(text: str) -> str:
    """PDF 字符串转义。"""
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_multipage_pdf(page_texts: list[str]) -> bytes:
    """手工构造一个多页 PDF（Helvetica，文本可被 pypdf 提取）。"""
    n = len(page_texts)
    font_num = 3
    page_nums: list[int] = []
    content_nums: list[int] = []
    cursor = 4
    for _ in range(n):
        page_nums.append(cursor)
        content_nums.append(cursor + 1)
        cursor += 2

    objects: dict[int, bytes] = {}
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{p} 0 R" for p in page_nums)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode("ascii")
    objects[font_num] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for index, text in enumerate(page_texts):
        page_num = page_nums[index]
        content_num = content_nums[index]
        objects[page_num] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_num} 0 R >> >> /Contents {content_num} 0 R >>"
        ).encode("ascii")
        y = 720
        ops = []
        for line in text.split("\n"):
            ops.append(f"BT /F1 12 Tf 72 {y} Td ({_pdf_escape(line)}) Tj ET")
            y -= 16
        content = "\n".join(ops).encode("latin-1", errors="replace")
        objects[content_num] = (
            b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n"
            + content + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for num in sorted(objects):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode("ascii") + objects[num] + b"\nendobj\n"
    xref_pos = len(out)
    max_num = max(objects)
    out += f"xref\n0 {max_num + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for num in range(1, max_num + 1):
        if num in offsets:
            out += f"{offsets[num]:010d} 00000 n \n".encode("ascii")
        else:
            out += b"0000000000 65535 f \n"
    out += (
        f"trailer\n<< /Size {max_num + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def build_docx() -> bytes:
    """生成含标题层级的 DOCX。"""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("第三章 导数与微分", level=1)
    doc.add_paragraph("导数的定义刻画了函数在一点处的瞬时变化率。")
    doc.add_heading("3.2 洛必达法则", level=2)
    doc.add_paragraph("洛必达法则是求未定式极限的重要方法，适用于 0/0 与 ∞/∞ 型。")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pptx() -> bytes:
    """生成 ≥3 页的 PPTX。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]
    slides_content = [
        ("第一讲 极限", "极限描述数列或函数在趋近过程中的变化趋势。"),
        ("第二讲 洛必达法则", "洛必达法则用于计算未定式极限。"),
        ("第三讲 习题", "练习：求 lim(x->0) sin(x)/x。"),
    ]
    for title, body in slides_content:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def build_scanned_pdf() -> bytes:
    """生成无文本层的图片 PDF（模拟扫描版）。"""
    from PIL import Image

    img = Image.new("RGB", (600, 400), color=(230, 230, 230))
    buffer = io.BytesIO()
    img.save(buffer, format="PDF")
    return buffer.getvalue()


def make_fixtures() -> dict[str, bytes]:
    """生成并落盘全部夹具，返回 {文件名: 字节}。"""
    OUT.mkdir(parents=True, exist_ok=True)
    fixtures: dict[str, bytes] = {}

    pdf_pages = [f"Page {i} - limit and derivative notes. Section {i}." for i in range(1, 6)]
    fixtures["高数-第3章.pdf"] = build_multipage_pdf(pdf_pages)
    fixtures["扫描版示例.pdf"] = build_scanned_pdf()
    fixtures["讲义.docx"] = build_docx()
    fixtures["课件.pptx"] = build_pptx()
    fixtures["笔记.md"] = "# 极限笔记\n\n## 洛必达法则\n\n洛必达法则适用于 0/0 型极限。\n\n## 常用技巧\n\n先检查是否可导。".encode("utf-8")
    fixtures["摘要.txt"] = "这是一份测试文本，用于验证 TXT 解析与检索。极限与连续。".encode("utf-8")

    for name, data in fixtures.items():
        (OUT / name).write_bytes(data)
    return fixtures


# ── 校验逻辑 ────────────────────────────────────────────
def upload_all(client: httpx.Client, fixtures: dict[str, bytes]) -> dict[str, str]:
    """批量上传并返回 {文件名: doc_id}（对重复文件复用库中已有条目，保证可重复运行）。"""
    files = [
        ("files", (name, data, "application/octet-stream"))
        for name, data in fixtures.items()
    ]
    resp = client.post("/api/documents/upload", files=files)
    body = resp.json()
    check("POST /documents/upload code=0", body.get("code") == 0, str(body.get("message")))
    ids = {d["title"]: d["id"] for d in body["data"]["documents"]}

    # 重复文件被跳过 → 从资料库按标题补齐 id。
    if len(ids) < len(fixtures):
        listed = client.get("/api/documents", params={"page_size": 200}).json()["data"]["items"]
        by_title = {d["title"]: d["id"] for d in listed}
        for title in fixtures:
            if title not in ids and title in by_title:
                ids[title] = by_title[title]

    check("全部文件均可用（新建或复用）", len(ids) == len(fixtures),
          f"{len(ids)}/{len(fixtures)} skipped={body['data'].get('skipped')}")
    return ids


def wait_ready(client: httpx.Client, doc_id: str, timeout: float = 60.0) -> dict:
    """轮询直到文档解析完成。"""
    deadline = time.time() + timeout
    last = {}
    while time.time() < deadline:
        last = client.get(f"/api/documents/{doc_id}").json()["data"]["document"]
        if last["status"] in ("ready", "failed"):
            return last
        time.sleep(0.4)
    return last


def _chunks(doc_id: str) -> list[sqlite3.Row]:
    """读取某文档的切片行。"""
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    try:
        return conn.execute(
            "SELECT ordinal, page_no, section FROM chunks WHERE document_id=? ORDER BY ordinal",
            (doc_id,),
        ).fetchall()
    finally:
        conn.close()


def _count_chunks(doc_id: str) -> int:
    """统计某文档切片数。"""
    conn = sqlite3.connect(str(DB_PATH))
    try:
        return int(conn.execute(
            "SELECT COUNT(*) FROM chunks WHERE document_id=?", (doc_id,)
        ).fetchone()[0])
    finally:
        conn.close()


def main() -> int:
    fixtures = make_fixtures()
    print(f"夹具已生成于 {OUT}")
    client = httpx.Client(base_url=BASE, timeout=90.0)

    ids = upload_all(client, fixtures)

    # 逐份等待 ready。
    for title in fixtures:
        doc = wait_ready(client, ids[title])
        if title == "扫描版示例.pdf":
            check(f"{title} status=ready（扫描版）", doc["status"] == "ready", doc["status"])
            check(f"{title} 带 warning 提示", bool(doc.get("warning")), str(doc.get("warning")))
            check(f"{title} 无可检索切片", _count_chunks(doc["id"]) == 0)
        else:
            check(f"{title} status=ready", doc["status"] == "ready", str(doc.get("error")))

    # PDF：切片 page_no 必须为整数且落在 1..5。
    pdf_rows = _chunks(ids["高数-第3章.pdf"])
    check("PDF 有切片", len(pdf_rows) >= 1, str(len(pdf_rows)))
    pdf_pages = [r["page_no"] for r in pdf_rows]
    check("PDF 切片 page_no 全为整数", all(isinstance(p, int) for p in pdf_pages), str(pdf_pages))
    check("PDF 切片 page_no 落在 1..5", all(p is not None and 1 <= p <= 5 for p in pdf_pages), str(pdf_pages))

    # PPTX：page_no 必须等于幻灯片序号（1..3）。
    pptx_doc = client.get(f"/api/documents/{ids['课件.pptx']}").json()["data"]["document"]
    check("PPTX page_count=3", pptx_doc["page_count"] == 3, str(pptx_doc["page_count"]))
    pptx_pages = [r["page_no"] for r in _chunks(ids["课件.pptx"])]
    check("PPTX 切片 page_no 为幻灯片序号(1..3)", all(p in (1, 2, 3) for p in pptx_pages), str(pptx_pages))

    # DOCX / MD 有 section。
    docx_sections = [r["section"] for r in _chunks(ids["讲义.docx"]) if r["section"]]
    check("DOCX 切片带 section 层级", len(docx_sections) >= 1, str(docx_sections[:2]))

    # 删除文档 → chunks 归零。
    target = ids["摘要.txt"]
    check("删除前 chunks>0", _count_chunks(target) > 0, str(_count_chunks(target)))
    resp = client.delete(f"/api/documents/{target}")
    check("DELETE /documents/{id} code=0", resp.json().get("code") == 0)
    check("删除后 chunks 归零", _count_chunks(target) == 0)

    print("-" * 60)
    print(f"全部通过：{PASSED} 项")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
